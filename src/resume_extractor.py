from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


def extract_docx(file_path):
    """Extract text from DOCX paragraphs and tables."""
    document = Document(file_path)
    text = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())

    # Extract tables.
    for table in document.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_data.append(cell_text)

            if row_data:
                text.append(" | ".join(row_data))

    return "\n".join(text)


def extract_pdf(file_path):
    """Extract text from PDF pages."""
    reader = PdfReader(file_path)
    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text and page_text.strip():
            text.append(page_text.strip())

    return "\n".join(text)


def extract_text(file_path):
    """Extract text from a supported resume file."""
    file_path = Path(file_path)

    if file_path.suffix.lower() == ".docx":
        return extract_docx(file_path)

    elif file_path.suffix.lower() == ".pdf":
        return extract_pdf(file_path)

    else:
        raise ValueError(
            f"Unsupported file format: {file_path.suffix}"
        )


def clean_text(text):
    """Remove excess spaces and empty lines."""
    text = re.sub(r"[ \t]+", " ", text)

    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]

    return "\n".join(lines)


def normalize_bullets(text):
    """Convert common bullet symbols and dash bullets into hyphens."""

    # Normalize common bullet symbols, including the white circle.
    text = re.sub(
        r"^[\u2022\u25AA\u25E6\u25CB\u25CF*]\s*",
        "- ",
        text,
        flags=re.MULTILINE
    )

    # Normalize lines beginning with en dashes or em dashes.
    text = re.sub(
        r"^[\u2013\u2014]\s*",
        "- ",
        text,
        flags=re.MULTILINE
    )

    return text


def normalize_section_headings(text):
    """Normalize common resume section headings."""

    section_headings = {
        "PROFILE": "PROFESSIONAL SUMMARY",
        "SUMMARY": "PROFESSIONAL SUMMARY",
        "PROFESSIONAL SUMMARY": "PROFESSIONAL SUMMARY",
        "WORK EXPERIENCE": "PROFESSIONAL EXPERIENCE",
        "EXPERIENCE": "PROFESSIONAL EXPERIENCE",
        "PROFESSIONAL EXPERIENCE": "PROFESSIONAL EXPERIENCE",
        "SKILLS": "SKILLS",
        "TECHNICAL SKILLS": "TECHNICAL SKILLS",
        "EDUCATION": "EDUCATION",
        "CERTIFICATIONS": "CERTIFICATIONS",
        "PROJECTS": "PROJECTS",
        "KEY PROJECT": "PROJECTS"
    }

    lines = []

    for line in text.splitlines():
        cleaned_line = line.strip()
        heading = section_headings.get(cleaned_line.upper())

        if heading:
            lines.append(heading)
        else:
            lines.append(cleaned_line)

    return "\n".join(lines)


def normalize_special_characters(text):
    """Normalize common Unicode characters used in resumes."""

    # Replace non-breaking spaces.
    text = text.replace("\u00a0", " ")

    # Normalize different dash characters.
    text = re.sub(
        r"[\u2010\u2011\u2012\u2013\u2014\u2015\u2212]",
        "-",
        text
    )

    # Normalize curly quotation marks.
    text = text.replace("“", '"')
    text = text.replace("”", '"')
    text = text.replace("‘", "'")
    text = text.replace("’", "'")

    return text
